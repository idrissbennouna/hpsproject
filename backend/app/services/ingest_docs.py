# ingest_docs.py
import sys
import os
from pathlib import Path

print("🔄 Démarrage du script d'ingestion...")

try:
    from dotenv import load_dotenv
    from langchain_community.vectorstores import PGVector
    from langchain_google_genai import GoogleGenerativeAIEmbeddings
    from langchain_core.documents import Document
    import openpyxl
    print("✅ Tous les modules requis ont été importés avec succès.")
except ImportError as e:
    print(f"❌ Erreur d'importation : {e}")
    print("👉 Installe les dépendances manquantes avec : pip install langchain-community langchain-google-genai pgvector psycopg2-binary openpyxl")
    sys.exit(1)

load_dotenv()


def charger_documents_depuis_excel(excel_path: Path) -> list[Document]:
    """Convertit chaque fonction documentée de Spec_PowerCARD.xlsx en Document réel (aucune donnée inventée)."""
    if not excel_path.exists():
        print(f"⚠️ Fichier Excel introuvable : {excel_path} — ignoré.")
        return []

    wb = openpyxl.load_workbook(str(excel_path), data_only=True)
    ws = wb["Lib"]
    documents = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        func_name, source, path, description, exception = row[:5]
        if func_name:
            content = (
                f"Fonction {func_name.strip()} (module {source}, fichier {path})\n"
                f"Description : {description}\n"
                f"Conditions : {exception}"
            )
            documents.append(Document(
                page_content=content,
                metadata={"source": "Spec_PowerCARD.xlsx", "function": func_name.strip()}
            ))
    print(f"📊 {len(documents)} fonctions extraites de Spec_PowerCARD.xlsx.")
    return documents


def initialiser_et_ingerer():
    CONNECTION_STRING = "postgresql+psycopg2://postgres:password@localhost:5432/hps_docs_db"
    COLLECTION_NAME = "hps_specifications"

    api_key = os.getenv("GOOGLE_API_KEY")
    if not api_key:
        print("❌ Erreur : La variable d'environnement GOOGLE_API_KEY est introuvable dans ton fichier .env")
        return

    try:
        embeddings = GoogleGenerativeAIEmbeddings(
            model="models/gemini-embedding-001",
            google_api_key=api_key
        )

        current_dir = Path(__file__).resolve().parent
        storage_dir = current_dir / "app" / "storage"
        excel_path = storage_dir / "Spec_PowerCARD.xlsx"

        documents_specs = charger_documents_depuis_excel(excel_path)

        if not documents_specs:
            print("❌ Aucun document réel à ingérer — vérifie le chemin du fichier Excel. Arrêt.")
            return

        print("📥 Connexion à PostgreSQL et initialisation de pgvector...")

        db = PGVector.from_documents(
            documents=documents_specs,
            embedding=embeddings,
            collection_name=COLLECTION_NAME,
            connection_string=CONNECTION_STRING,
            pre_delete_collection=True
        )

        print(f"✅ Base de données PostgreSQL initialisée avec succès : {len(documents_specs)} spécifications réelles indexées !")

    except Exception as e:
        print(f"❌ Une erreur est survenue lors de l'initialisation de la DB : {e}")


if __name__ == "__main__":
    initialiser_et_ingerer()